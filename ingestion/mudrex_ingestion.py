"""Read and save raw Mudrex transaction data."""

import argparse
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from storage.local_storage import LocalStorage
from storage.s3_storage import S3Storage


REPORT_FILE = (
    Path("raw report")
    / "mudrex"
    / "4801485_transaction_statement.xlsx"
)


def read_sheet(
    sheet_name,
    header_row
):
    """Read one Excel sheet and return its data rows."""

    workbook = load_workbook(
        REPORT_FILE,
        read_only=True,
        data_only=True
    )

    worksheet = workbook[sheet_name]

    rows = worksheet.iter_rows(
        min_row=header_row,
        values_only=True
    )

    headers = next(rows)

    records = []

    for row in rows:

        if not any(
            value is not None
            for value in row
        ):
            continue

        record = dict(
            zip(headers, row)
        )

        records.append(record)

    workbook.close()

    return records


def get_usdt_spot_trades():
    """Read historical USDT Spot trades."""

    return read_sheet(
        sheet_name="USDT-Spot",
        header_row=5
    )


def get_inr_spot_trades():
    """Read newer INR Spot trades."""

    return read_sheet(
        sheet_name="INR-Spot",
        header_row=7
    )


def get_deposits():
    """Read fiat and crypto deposit history."""

    return read_sheet(
        sheet_name="deposit_history",
        header_row=8
    )


def get_withdrawals():
    """Read fiat and crypto withdrawal history."""

    return read_sheet(
        sheet_name="withdrawal_history",
        header_row=9
    )


def get_earn_history():
    """Read Mudrex Earn transfers and interest rewards."""

    return read_sheet(
        sheet_name="Earn history",
        header_row=1
    )


def get_usdt_fund_transactions():
    """Read INR-to-USDT and USDT-to-INR transactions."""

    return read_sheet(
        sheet_name="USDT Funds History",
        header_row=7
    )


def save_raw_data(
    data,
    local_storage,
    s3_storage,
    data_type
):
    """Save Mudrex data locally and optionally upload it to S3."""

    file_path = local_storage.save_json(
        data=data,
        source="mudrex",
        data_type=data_type
    )

    if s3_storage is not None:

        s3_storage.upload_file(
            local_file=file_path,
            s3_key=file_path.as_posix()
        )

    print(
        f"Saved {data_type}: {file_path}"
    )

    return file_path


def run_ingestion(
    local_storage=None,
    s3_storage=None,
    upload_to_s3=True
):
    """Read and save all useful Mudrex datasets."""

    local_storage = (
        local_storage
        or LocalStorage()
    )

    if upload_to_s3:

        s3_storage = (
            s3_storage
            or S3Storage()
        )

    else:

        s3_storage = None

    saved_files = {}

    def save(
        name,
        data
    ):

        saved_files[name] = save_raw_data(
            data=data,
            local_storage=local_storage,
            s3_storage=s3_storage,
            data_type=name
        )

    print(
        "1. Reading Mudrex USDT Spot trades..."
    )

    save(
        name="usdt_spot_trades",
        data=get_usdt_spot_trades()
    )

    print(
        "2. Reading Mudrex INR Spot trades..."
    )

    save(
        name="inr_spot_trades",
        data=get_inr_spot_trades()
    )

    print(
        "3. Reading Mudrex deposits..."
    )

    save(
        name="deposits",
        data=get_deposits()
    )

    print(
        "4. Reading Mudrex withdrawals..."
    )

    save(
        name="withdrawals",
        data=get_withdrawals()
    )

    print(
        "5. Reading Mudrex Earn history..."
    )

    save(
        name="earn_history",
        data=get_earn_history()
    )

    print(
        "6. Reading Mudrex USDT fund transactions..."
    )

    save(
        name="usdt_fund_transactions",
        data=get_usdt_fund_transactions()
    )

    print(
        "7. Creating ingestion manifest..."
    )

    manifest = {
        "source": "mudrex",
        "source_report": str(REPORT_FILE),
        "completed_at": datetime.now(
            timezone.utc
        ).isoformat(),
        "files": {
            name: str(path)
            for name, path in saved_files.items()
        }
    }

    save(
        name="ingestion_manifest",
        data=manifest
    )

    return saved_files


def main():
    """Read terminal options and start Mudrex ingestion."""

    parser = argparse.ArgumentParser(
        description="Read and save raw Mudrex data"
    )

    parser.add_argument(
        "--local-only",
        "--no-s3",
        action="store_true",
        dest="local_only",
        help="Save locally without uploading to S3"
    )

    args = parser.parse_args()

    saved_files = run_ingestion(
        upload_to_s3=not args.local_only
    )

    print(
        f"Done. Saved {len(saved_files)} raw datasets."
    )


if __name__ == "__main__":
    main()